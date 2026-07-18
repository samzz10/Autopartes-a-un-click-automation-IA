"""
Conciliación Bancaria — Banco Colombia vs Alegra
=================================================
Módulo para agregar al bot de cartera (cartera_bot.py).

Flujo:
1. Usuario escribe /conciliacion
2. Bot pide el extracto PDF del Banco Colombia
3. Usuario lo adjunta
4. Bot extrae los movimientos del PDF
5. Bot consulta gastos e ingresos de Alegra
6. Compara ambos y envía Excel con resultado

Requisitos adicionales:
    pip install pdfplumber openpyxl
"""

import io
import re
import base64
import requests
import pdfplumber
from datetime import datetime, date
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from telegram import Update
from telegram.ext import ContextTypes, ConversationHandler, CommandHandler, MessageHandler, filters

# ══════════════════════════════════════════════════════
#  CONFIGURACIÓN
# ══════════════════════════════════════════════════════
ALEGRA_TOKEN = "0d9e3617130df4ec9b11"
ALEGRA_EMAIL = "autopartesaunclickcol@gmail.com"

# Estado de conversación
ESPERANDO_PDF = 1


# ══════════════════════════════════════════════════════
#  LEER PDF DEL BANCO COLOMBIA
# ══════════════════════════════════════════════════════
def leer_extracto_banco(pdf_bytes):
    """
    Extrae movimientos del extracto PDF del Banco Colombia usando texto.
    Formato de línea: DD/MM  DESCRIPCION  VALOR  SALDO
    """
    movimientos = []
    patron = re.compile(r'^(\d{1,2}/\d{2})\s+(.+?)\s+(-?[\d,]+\.\d{2})\s+(-?[\d,]+\.\d{2})$')
    anio = date.today().year

    with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
        for page in pdf.pages:
            text = page.extract_text()
            if not text:
                continue
            for line in text.split("\n"):
                line = line.strip()
                m = patron.match(line)
                if not m:
                    continue
                fecha_str, descripcion, valor_str, _ = m.groups()
                try:
                    dia, mes = fecha_str.split("/")
                    fecha = date(anio, int(mes), int(dia))
                except Exception:
                    continue
                valor = float(valor_str.replace(",", ""))
                movimientos.append({
                    "fecha":       fecha,
                    "descripcion": descripcion.strip(),
                    "valor":       valor,
                    "fuente":      "banco",
                })

    return movimientos


# ══════════════════════════════════════════════════════
#  CONSULTAR ALEGRA — GASTOS E INGRESOS
# ══════════════════════════════════════════════════════
def obtener_movimientos_alegra(fecha_desde, fecha_hasta):
    """
    Consulta pagos (ingresos) y gastos de Alegra en el período dado.
    Filtra manualmente por fecha ya que la API ignora los parámetros de fecha.
    """
    auth = base64.b64encode(f"{ALEGRA_EMAIL}:{ALEGRA_TOKEN}".encode()).decode()
    headers = {
        "Authorization": f"Basic {auth}",
        "Accept": "application/json"
    }
    base_url = "https://api.alegra.com/api/v1"
    movimientos = []

    # ── Ingresos: pagos recibidos ──────────────────────
    start = 0
    while True:
        resp = requests.get(f"{base_url}/payments", headers=headers, params={
            "start": start, "limit": 30,
            "order_field": "date", "order_direction": "DESC",
        }, timeout=120)
        if not resp.ok:
            break
        data = resp.json()
        if not data:
            break

        fechas_pagina = []
        for p in data:
            try:
                fecha = datetime.strptime(p.get("date", "")[:10], "%Y-%m-%d").date()
            except Exception:
                continue
            fechas_pagina.append(fecha)
            # Filtrar manualmente por rango
            if fecha < fecha_desde or fecha > fecha_hasta:
                continue
            movimientos.append({
                "fecha":       fecha,
                "descripcion": f"Pago - {p.get('client', {}).get('name', '') or p.get('observations', '')}",
                "valor":       float(p.get("amount", 0) or 0),
                "fuente":      "alegra",
            })

        # Si todos los registros de la página son anteriores al rango, parar
        if fechas_pagina and max(fechas_pagina) < fecha_desde:
            break
        if len(data) < 30:
            break
        start += 30

    # ── Gastos ────────────────────────────────────────
    start = 0
    while True:
        resp = requests.get(f"{base_url}/bills", headers=headers, params={
            "start": start, "limit": 30,
            "order_field": "date", "order_direction": "DESC",
        }, timeout=120)
        if not resp.ok:
            break
        data = resp.json()
        if not data:
            break

        fechas_pagina = []
        for g in data:
            try:
                fecha = datetime.strptime(g.get("date", "")[:10], "%Y-%m-%d").date()
            except Exception:
                continue
            fechas_pagina.append(fecha)
            # Filtrar manualmente por rango
            if fecha < fecha_desde or fecha > fecha_hasta:
                continue
            movimientos.append({
                "fecha":       fecha,
                "descripcion": f"Gasto - {g.get('provider', {}).get('name', '') or g.get('observations', '')}",
                "valor":       -float(g.get("total", 0) or 0),
                "fuente":      "alegra",
            })

        # Si todos los registros de la página son anteriores al rango, parar
        if fechas_pagina and max(fechas_pagina) < fecha_desde:
            break
        if len(data) < 30:
            break
        start += 30

    return movimientos


# ══════════════════════════════════════════════════════
#  COMPARAR Y GENERAR EXCEL
# ══════════════════════════════════════════════════════
def conciliar_y_generar_excel(banco, alegra):
    """
    Compara movimientos del banco vs Alegra.
    Criterio: mismo valor (±1000 tolerancia) y fecha cercana (±3 días).
    """
    coinciden   = []
    solo_banco  = []
    solo_alegra = []

    alegra_usados = [False] * len(alegra)

    for mov_b in banco:
        encontrado = False
        for j, mov_a in enumerate(alegra):
            if alegra_usados[j]:
                continue
            diff_valor = abs(float(mov_b["valor"]) - float(mov_a["valor"]))
            diff_dias  = abs((mov_b["fecha"] - mov_a["fecha"]).days)
            if diff_valor <= 1000 and diff_dias <= 3:
                coinciden.append({
                    "Fecha Banco":  mov_b["fecha"].strftime("%d/%m/%Y"),
                    "Desc. Banco":  mov_b["descripcion"],
                    "Valor Banco":  float(mov_b["valor"]),
                    "Fecha Alegra": mov_a["fecha"].strftime("%d/%m/%Y"),
                    "Desc. Alegra": mov_a["descripcion"],
                    "Valor Alegra": float(mov_a["valor"]),
                    "Diferencia":   round(float(mov_b["valor"]) - float(mov_a["valor"]), 2),
                })
                alegra_usados[j] = True
                encontrado = True
                break
        if not encontrado:
            solo_banco.append(mov_b)

    for j, mov_a in enumerate(alegra):
        if not alegra_usados[j]:
            solo_alegra.append(mov_a)

    # ── Estilos ───────────────────────────────────────
    C_NAVY    = "1A3A5C"
    C_BLUE    = "2E7DBA"
    C_WHITE   = "FFFFFF"
    C_GREEN   = "E8F5E9"
    C_GREEN2  = "2E7D32"
    C_RED     = "FFEBEE"
    C_RED2    = "C62828"
    C_ORANGE  = "FFF3E0"
    C_ORANGE2 = "E65100"
    C_BORDER  = "B8D0E8"

    def fill(h): return PatternFill("solid", fgColor=h)
    thin  = Side(style="thin",   color=C_BORDER)
    thick = Side(style="medium", color=C_BLUE)
    b  = Border(left=thin,  right=thin,  top=thin,  bottom=thin)
    bH = Border(left=thick, right=thick, top=thick, bottom=thick)

    wb = Workbook()

    # ── HOJA 1: RESUMEN ───────────────────────────────
    ws = wb.active
    ws.title = "Resumen"

    def sc(row, col, val, font=None, pfill=None, align=None, border=None, nfmt=None):
        c = ws.cell(row=row, column=col, value=val)
        if font:   c.font   = font
        if pfill:  c.fill   = pfill
        if align:  c.alignment = align
        if border: c.border = border
        if nfmt:   c.number_format = nfmt
        return c

    def ms(ref, val, font=None, pfill=None, align=None, border=None):
        c = ws[ref]
        c.value = val
        if font:   c.font   = font
        if pfill:  c.fill   = pfill
        if align:  c.alignment = align
        if border: c.border = border
        return c

    ws.merge_cells("A1:F1")
    ms("A1", "CONCILIACIÓN BANCARIA — BANCO COLOMBIA vs ALEGRA",
       Font(name="Arial", bold=True, size=14, color=C_WHITE),
       fill(C_NAVY), Alignment(horizontal="center", vertical="center"), bH)
    ws.row_dimensions[1].height = 32

    ws.merge_cells("A2:F2")
    ms("A2", f"Generado el {datetime.now().strftime('%d/%m/%Y %H:%M')}",
       Font(name="Arial", size=10, italic=True, color=C_WHITE),
       fill(C_BLUE), Alignment(horizontal="center", vertical="center"), b)
    ws.row_dimensions[2].height = 18
    ws.row_dimensions[3].height = 8

    kpis = [
        ("A4:B5", "✅ Coinciden",       str(len(coinciden)),   C_GREEN,  C_GREEN2),
        ("C4:D5", "❌ Solo en banco",    str(len(solo_banco)),  C_RED,    C_RED2),
        ("E4:F5", "⚠️ Solo en Alegra",  str(len(solo_alegra)), C_ORANGE, C_ORANGE2),
    ]
    for rng, label, value, bg, fg in kpis:
        r1 = rng.split(":")[0]
        r2 = rng.split(":")[1]
        ws.merge_cells(f"{r1[0]}4:{r2[0]}4")
        ws.merge_cells(f"{r1[0]}5:{r2[0]}5")
        ms(f"{r1[0]}4", label,
           Font(name="Arial", size=9, color=fg),
           fill(bg), Alignment(horizontal="center", vertical="bottom"), b)
        ms(f"{r1[0]}5", value,
           Font(name="Arial", bold=True, size=18, color=fg),
           fill(bg), Alignment(horizontal="center", vertical="top"), b)

    ws.row_dimensions[4].height = 18
    ws.row_dimensions[5].height = 32
    ws.row_dimensions[6].height = 8

    for col, w in zip("ABCDEF", [16, 35, 16, 16, 35, 16]):
        ws.column_dimensions[col].width = w

    # ── HOJA 2: COINCIDEN ─────────────────────────────
    ws2 = wb.create_sheet("Coinciden")
    hdrs = ["Fecha Banco","Desc. Banco","Valor Banco","Fecha Alegra","Desc. Alegra","Valor Alegra","Diferencia"]
    for col, h in enumerate(hdrs, 1):
        c = ws2.cell(row=1, column=col, value=h)
        c.font      = Font(name="Arial", bold=True, size=10, color=C_WHITE)
        c.fill      = fill(C_GREEN2)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border    = b
    ws2.row_dimensions[1].height = 22

    for i, row in enumerate(coinciden, 2):
        for col, key in enumerate(hdrs, 1):
            c = ws2.cell(row=i, column=col, value=row[key])
            c.fill      = fill(C_GREEN if i % 2 == 0 else "F1F8E9")
            c.border    = b
            c.font      = Font(name="Arial", size=9)
            c.alignment = Alignment(vertical="center", wrap_text=True,
                                    horizontal="right" if key in ("Valor Banco","Valor Alegra","Diferencia") else "left")
            if key in ("Valor Banco","Valor Alegra","Diferencia"):
                c.number_format = '"$ "#,##0'
        ws2.row_dimensions[i].height = 28

    for col, w in zip("ABCDEFG", [13,38,14,13,38,14,13]):
        ws2.column_dimensions[col].width = w

    # ── HOJA 3: SOLO EN BANCO ─────────────────────────
    ws3 = wb.create_sheet("Solo en banco")
    hdrs3 = ["Fecha","Descripción","Valor"]
    for col, h in enumerate(hdrs3, 1):
        c = ws3.cell(row=1, column=col, value=h)
        c.font      = Font(name="Arial", bold=True, size=10, color=C_WHITE)
        c.fill      = fill(C_RED2)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border    = b
    ws3.row_dimensions[1].height = 22

    for i, mov in enumerate(solo_banco, 2):
        ws3.cell(row=i, column=1, value=mov["fecha"].strftime("%d/%m/%Y")).border = b
        ws3.cell(row=i, column=2, value=mov["descripcion"]).border = b
        c = ws3.cell(row=i, column=3, value=float(mov["valor"]))
        c.number_format = '"$ "#,##0'
        c.border = b
        for col in range(1, 4):
            ws3.cell(row=i, column=col).fill      = fill(C_RED if i % 2 == 0 else "FFCDD2")
            ws3.cell(row=i, column=col).font      = Font(name="Arial", size=9)
            ws3.cell(row=i, column=col).alignment = Alignment(vertical="center", wrap_text=True)
        ws3.row_dimensions[i].height = 28

    ws3.column_dimensions["A"].width = 13
    ws3.column_dimensions["B"].width = 50
    ws3.column_dimensions["C"].width = 16

    # ── HOJA 4: SOLO EN ALEGRA ────────────────────────
    ws4 = wb.create_sheet("Solo en Alegra")
    for col, h in enumerate(hdrs3, 1):
        c = ws4.cell(row=1, column=col, value=h)
        c.font      = Font(name="Arial", bold=True, size=10, color=C_WHITE)
        c.fill      = fill(C_ORANGE2)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border    = b
    ws4.row_dimensions[1].height = 22

    for i, mov in enumerate(solo_alegra, 2):
        ws4.cell(row=i, column=1, value=mov["fecha"].strftime("%d/%m/%Y")).border = b
        ws4.cell(row=i, column=2, value=mov["descripcion"]).border = b
        c = ws4.cell(row=i, column=3, value=float(mov["valor"]))
        c.number_format = '"$ "#,##0'
        c.border = b
        for col in range(1, 4):
            ws4.cell(row=i, column=col).fill      = fill(C_ORANGE if i % 2 == 0 else "FFE0B2")
            ws4.cell(row=i, column=col).font      = Font(name="Arial", size=9)
            ws4.cell(row=i, column=col).alignment = Alignment(vertical="center", wrap_text=True)
        ws4.row_dimensions[i].height = 28

    ws4.column_dimensions["A"].width = 13
    ws4.column_dimensions["B"].width = 50
    ws4.column_dimensions["C"].width = 16

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read(), len(coinciden), len(solo_banco), len(solo_alegra)


# ══════════════════════════════════════════════════════
#  HANDLERS DEL BOT
# ══════════════════════════════════════════════════════
async def conciliacion_start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await update.message.reply_text(
        "🏦 *Conciliación Bancaria*\n\n"
        "Envíame el extracto del *Banco Colombia* en PDF y lo comparo automáticamente con Alegra.\n\n"
        "📎 Adjunta el archivo PDF ahora:",
        parse_mode="Markdown"
    )
    return ESPERANDO_PDF


async def conciliacion_recibir_pdf(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not update.message.document:
        await update.message.reply_text("❌ Por favor adjunta un archivo PDF.")
        return ESPERANDO_PDF

    doc = update.message.document
    if not doc.file_name.lower().endswith(".pdf"):
        await update.message.reply_text("❌ El archivo debe ser PDF. Inténtalo de nuevo.")
        return ESPERANDO_PDF

    await update.message.reply_text("⏳ Procesando... esto puede tardar hasta 2 minutos, ten paciencia.")

    try:
        file = await context.bot.get_file(doc.file_id)
        pdf_bytes = bytes(await file.download_as_bytearray())

        await update.message.reply_text("📄 Leyendo extracto del banco...")
        banco = leer_extracto_banco(pdf_bytes)

        if not banco:
            await update.message.reply_text(
                "❌ No pude leer movimientos del PDF. Verifica que sea el extracto del Banco Colombia."
            )
            return ConversationHandler.END

        fechas  = [m["fecha"] for m in banco]
        f_desde = min(fechas)
        f_hasta = max(fechas)

        await update.message.reply_text(
            f"✅ {len(banco)} movimientos encontrados en el banco\n"
            f"📅 Período: {f_desde.strftime('%d/%m/%Y')} — {f_hasta.strftime('%d/%m/%Y')}\n\n"
            f"🔄 Consultando Alegra..."
        )

        alegra = obtener_movimientos_alegra(f_desde, f_hasta)
        await update.message.reply_text(
            f"✅ {len(alegra)} movimientos encontrados en Alegra\n\n🔀 Conciliando..."
        )

        excel_bytes, coinciden, solo_banco, solo_alegra = conciliar_y_generar_excel(banco, alegra)

        resumen = (
            f"📊 *Resultado de Conciliación*\n"
            f"📅 {f_desde.strftime('%d/%m/%Y')} — {f_hasta.strftime('%d/%m/%Y')}\n\n"
            f"✅ Coinciden: *{coinciden}*\n"
            f"❌ Solo en banco (no en Alegra): *{solo_banco}*\n"
            f"⚠️ Solo en Alegra (no en banco): *{solo_alegra}*\n\n"
        )
        if solo_banco > 0 or solo_alegra > 0:
            resumen += "⚠️ *Hay diferencias que requieren revisión.* Ver hojas del Excel."
        else:
            resumen += "🎉 *Todo cuadra perfectamente.*"

        await update.message.reply_text(resumen, parse_mode="Markdown")

        nombre = f"Conciliacion_{f_desde.strftime('%Y%m%d')}_{f_hasta.strftime('%Y%m%d')}.xlsx"
        await update.message.reply_document(
            document=excel_bytes,
            filename=nombre,
            caption="📋 Conciliación — 4 hojas: Resumen, Coinciden, Solo banco, Solo Alegra"
        )

    except Exception as e:
        await update.message.reply_text(f"❌ Error durante la conciliación: {e}")

    return ConversationHandler.END


async def conciliacion_cancelar(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await update.message.reply_text("❌ Conciliación cancelada.")
    return ConversationHandler.END


def get_conciliacion_handler():
    return ConversationHandler(
        entry_points=[CommandHandler("conciliacion", conciliacion_start)],
        states={
            ESPERANDO_PDF: [
                MessageHandler(filters.Document.ALL, conciliacion_recibir_pdf)
            ],
        },
        fallbacks=[CommandHandler("cancelar", conciliacion_cancelar)],
    )